"""Unit tests for XLSX, DOCX, TXT, and MD loaders plus DocumentNormalizer."""

from __future__ import annotations

from pathlib import Path

import docx
import openpyxl
import pytest

from pygramattic_reports.exceptions import LoaderError, NormalizationError
from pygramattic_reports.loaders import (
    DocxLoader,
    MdLoader,
    TxtLoader,
    XlsxLoader,
    create_default_registry,
)
from pygramattic_reports.models import (
    ContentType,
    FileSourceConfig,
    RawData,
    SourceConfig,
    SourceType,
)
from pygramattic_reports.models.base import generate_id, now_utc
from pygramattic_reports.normalizers import DocumentNormalizer

FIXTURES_DIR = Path(__file__).parent.parent / "fixtures"


# ---- Fixture helpers ----


def _xlsx_config(path: Path, **kwargs: object) -> SourceConfig:
    return SourceConfig(
        source_type=SourceType.XLSX,
        name="test_xlsx",
        file=FileSourceConfig(path=path, **kwargs),
    )


def _docx_config(path: Path) -> SourceConfig:
    return SourceConfig(
        source_type=SourceType.DOCX,
        name="test_docx",
        file=FileSourceConfig(path=path),
    )


def _txt_config(path: Path, **kwargs: object) -> SourceConfig:
    return SourceConfig(
        source_type=SourceType.TXT,
        name="test_txt",
        file=FileSourceConfig(path=path, **kwargs),
    )


def _md_config(path: Path) -> SourceConfig:
    return SourceConfig(
        source_type=SourceType.MD,
        name="test_md",
        file=FileSourceConfig(path=path),
    )


@pytest.fixture
def sample_xlsx(tmp_path):
    """Create a simple XLSX workbook programmatically."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Revenue"
    ws.append(["region", "revenue", "quarter"])
    ws.append(["US", 1500.50, "Q1"])
    ws.append(["EU", 2300.75, "Q1"])
    ws.append(["APAC", 890.25, "Q1"])
    path = tmp_path / "test.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def multi_sheet_xlsx(tmp_path):
    """Create an XLSX workbook with multiple sheets."""
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Sales"
    ws1.append(["product", "units"])
    ws1.append(["Widget", 100])

    ws2 = wb.create_sheet("Costs")
    ws2.append(["category", "amount"])
    ws2.append(["Materials", 5000])
    ws2.append(["Labor", 8000])

    path = tmp_path / "multi.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def merged_xlsx(tmp_path):
    """Create an XLSX workbook with merged cells."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["region", "revenue"])
    ws.merge_cells("A2:A3")
    ws.cell(row=2, column=1, value="US")
    ws.cell(row=2, column=2, value=1000)
    ws.cell(row=3, column=2, value=2000)
    path = tmp_path / "merged.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def sample_docx(tmp_path):
    """Create a simple DOCX document programmatically."""
    doc = docx.Document()
    doc.add_heading("Executive Summary", level=1)
    doc.add_paragraph("Revenue grew 12% year-over-year.")
    doc.add_heading("Details", level=2)
    doc.add_paragraph("US market led with $1.5M.")

    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Region"
    table.cell(0, 1).text = "Revenue"
    table.cell(1, 0).text = "US"
    table.cell(1, 1).text = "1500"
    table.cell(2, 0).text = "EU"
    table.cell(2, 1).text = "2300"

    path = tmp_path / "test.docx"
    doc.save(path)
    return path


# ---- XLSX Loader ----


class TestXlsxLoaderBasic:
    def test_loads_single_sheet(self, sample_xlsx):
        loader = XlsxLoader()
        config = _xlsx_config(sample_xlsx)
        raw = loader.load(config)

        assert raw.content_type == ContentType.TABULAR
        assert raw.tabular_headers == ["region", "revenue", "quarter"]
        assert raw.row_count == 3
        assert raw.tabular_data[0]["region"] == "US"
        assert raw.source_checksum is not None

    def test_supported_types(self):
        loader = XlsxLoader()
        assert SourceType.XLSX in loader.supported_types()


class TestXlsxLoaderSpecificSheet:
    def test_loads_named_sheet(self, multi_sheet_xlsx):
        loader = XlsxLoader()
        config = _xlsx_config(multi_sheet_xlsx, sheet_name="Costs")
        raw = loader.load(config)

        assert raw.tabular_headers == ["category", "amount"]
        assert raw.row_count == 2
        assert raw.tabular_data[0]["category"] == "Materials"

    def test_missing_sheet_raises(self, multi_sheet_xlsx):
        loader = XlsxLoader()
        config = _xlsx_config(multi_sheet_xlsx, sheet_name="Nonexistent")

        with pytest.raises(LoaderError, match="not found"):
            loader.load(config)


class TestXlsxLoaderNoHeader:
    def test_auto_generates_column_names(self, tmp_path):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append([10, 20, 30])
        ws.append([40, 50, 60])
        path = tmp_path / "no_header.xlsx"
        wb.save(path)

        loader = XlsxLoader()
        config = _xlsx_config(path, has_header=False)
        raw = loader.load(config)

        assert raw.tabular_headers == ["col_0", "col_1", "col_2"]
        assert raw.row_count == 2

    def test_skips_empty_rows(self, tmp_path):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["a", "b"])
        ws.append([1, 2])
        ws.append([None, None])
        ws.append([3, 4])
        path = tmp_path / "empty_rows.xlsx"
        wb.save(path)

        loader = XlsxLoader()
        config = _xlsx_config(path)
        raw = loader.load(config)

        assert raw.row_count == 2


class TestXlsxLoaderMergedCells:
    def test_fills_merged_values(self, merged_xlsx):
        loader = XlsxLoader()
        config = _xlsx_config(merged_xlsx)
        raw = loader.load(config)

        assert raw.row_count == 2
        assert raw.tabular_data[0]["region"] == "US"
        assert raw.tabular_data[1]["region"] == "US"


class TestXlsxLoaderFileNotFound:
    def test_raises_loader_error(self, tmp_path):
        loader = XlsxLoader()
        config = _xlsx_config(tmp_path / "nonexistent.xlsx")

        with pytest.raises(LoaderError, match="File not found"):
            loader.load(config)


# ---- DOCX Loader ----


class TestDocxLoaderParagraphs:
    def test_extracts_paragraphs(self, sample_docx):
        loader = DocxLoader()
        config = _docx_config(sample_docx)
        raw = loader.load(config)

        assert raw.content_type == ContentType.DOCUMENT
        assert raw.document_text is not None
        assert len(raw.document_text) == 4

        # Headings should have # prefixes
        assert raw.document_text[0] == "# Executive Summary"
        assert raw.document_text[1] == "Revenue grew 12% year-over-year."
        assert raw.document_text[2] == "## Details"
        assert raw.document_text[3] == "US market led with $1.5M."

    def test_supported_types(self):
        loader = DocxLoader()
        assert SourceType.DOCX in loader.supported_types()


class TestDocxLoaderWithTable:
    def test_extracts_first_table(self, sample_docx):
        loader = DocxLoader()
        config = _docx_config(sample_docx)
        raw = loader.load(config)

        assert raw.tabular_headers == ["Region", "Revenue"]
        assert raw.tabular_data is not None
        assert len(raw.tabular_data) == 2
        assert raw.tabular_data[0]["Region"] == "US"
        assert raw.tabular_data[0]["Revenue"] == "1500"

    def test_file_not_found_raises(self, tmp_path):
        loader = DocxLoader()
        config = _docx_config(tmp_path / "nonexistent.docx")

        with pytest.raises(LoaderError, match="File not found"):
            loader.load(config)


# ---- TXT Loader ----


class TestTxtLoaderBasic:
    def test_loads_plain_text(self):
        loader = TxtLoader()
        config = _txt_config(FIXTURES_DIR / "sample.txt")
        raw = loader.load(config)

        assert raw.content_type == ContentType.DOCUMENT
        assert raw.document_text is not None
        assert len(raw.document_text) == 6
        assert raw.document_text[0] == "Executive Summary"
        assert raw.source_checksum is not None

    def test_supported_types(self):
        loader = TxtLoader()
        assert SourceType.TXT in loader.supported_types()


class TestTxtLoaderEncoding:
    def test_handles_latin1(self, tmp_path):
        txt_file = tmp_path / "latin1.txt"
        txt_file.write_bytes("Z\xfcrich is nice.\n\nGreat city.".encode("latin-1"))

        loader = TxtLoader()
        config = _txt_config(txt_file, encoding="latin-1")
        raw = loader.load(config)

        assert raw.document_text is not None
        assert raw.document_text[0] == "Z\u00fcrich is nice."

    def test_wrong_encoding_raises(self, tmp_path):
        txt_file = tmp_path / "bad.txt"
        txt_file.write_bytes(b"\xff\xfe\x00\x00invalid")

        loader = TxtLoader()
        config = _txt_config(txt_file)

        with pytest.raises(LoaderError, match="Encoding error"):
            loader.load(config)

    def test_file_not_found_raises(self, tmp_path):
        loader = TxtLoader()
        config = _txt_config(tmp_path / "nonexistent.txt")

        with pytest.raises(LoaderError, match="File not found"):
            loader.load(config)


# ---- MD Loader ----


class TestMdLoaderStructure:
    def test_parses_headings_and_paragraphs(self):
        loader = MdLoader()
        config = _md_config(FIXTURES_DIR / "sample.md")
        raw = loader.load(config)

        assert raw.content_type == ContentType.DOCUMENT
        assert raw.document_text is not None

        # Should have headings with # prefixes
        assert any(block.startswith("# ") for block in raw.document_text)
        assert any(block.startswith("## ") for block in raw.document_text)
        assert any("Revenue grew" in block for block in raw.document_text)

    def test_extracts_code_blocks(self):
        loader = MdLoader()
        config = _md_config(FIXTURES_DIR / "sample.md")
        raw = loader.load(config)

        code_blocks = [b for b in raw.document_text if b.startswith("```")]
        assert len(code_blocks) == 1
        assert "total = sum" in code_blocks[0]

    def test_extracts_lists(self):
        loader = MdLoader()
        config = _md_config(FIXTURES_DIR / "sample.md")
        raw = loader.load(config)

        list_blocks = [b for b in raw.document_text if b.startswith("- ")]
        assert len(list_blocks) == 1
        assert "US market leads" in list_blocks[0]

    def test_supported_types(self):
        loader = MdLoader()
        assert SourceType.MD in loader.supported_types()


class TestMdLoaderWithTable:
    def test_extracts_table(self):
        loader = MdLoader()
        config = _md_config(FIXTURES_DIR / "sample.md")
        raw = loader.load(config)

        assert raw.tabular_headers is not None
        assert "Region" in raw.tabular_headers
        assert raw.tabular_data is not None
        assert len(raw.tabular_data) == 3
        assert raw.tabular_data[0]["Region"] == "US"

    def test_file_not_found_raises(self, tmp_path):
        loader = MdLoader()
        config = _md_config(tmp_path / "nonexistent.md")

        with pytest.raises(LoaderError, match="File not found"):
            loader.load(config)


class TestMdLoaderFrontmatter:
    def test_strips_yaml_frontmatter(self, tmp_path):
        md_file = tmp_path / "frontmatter.md"
        md_file.write_text("---\ntitle: Test\nauthor: Bot\n---\n# Hello\n\nWorld\n")

        loader = MdLoader()
        config = _md_config(md_file)
        raw = loader.load(config)

        assert raw.document_text is not None
        # Frontmatter should be stripped
        assert not any("title:" in block for block in raw.document_text)
        assert any(block == "# Hello" for block in raw.document_text)


# ---- Document Normalizer ----


class TestDocumentNormalizer:
    def test_normalizes_document(self):
        raw = RawData(
            id=generate_id(),
            source_config=SourceConfig(
                source_type=SourceType.TXT,
                name="test_doc",
                file=FileSourceConfig(path=Path("/test/doc.txt")),
            ),
            content_type=ContentType.DOCUMENT,
            document_text=[
                "# Executive Summary",
                "Revenue grew 12%.",
                "## Details",
                "US market led.",
            ],
            loaded_at=now_utc(),
        )
        normalizer = DocumentNormalizer()
        dataset = normalizer.normalize(raw)

        assert dataset.row_count == 4
        assert list(dataset.dataframe.columns) == [
            "section_index",
            "section_type",
            "level",
            "content",
        ]

        types = dataset.dataframe["section_type"].tolist()
        assert types[0] == "heading"
        assert types[1] == "paragraph"
        assert types[2] == "heading"
        assert types[3] == "paragraph"

        levels = dataset.dataframe["level"].tolist()
        assert levels[0] == 1
        assert levels[2] == 2

        contents = dataset.dataframe["content"].tolist()
        assert contents[0] == "Executive Summary"
        assert contents[1] == "Revenue grew 12%."

    def test_normalizes_code_block(self):
        raw = RawData(
            id=generate_id(),
            source_config=SourceConfig(
                source_type=SourceType.MD,
                name="code_doc",
                file=FileSourceConfig(path=Path("/test/code.md")),
            ),
            content_type=ContentType.DOCUMENT,
            document_text=["```python\nprint('hello')\n```"],
            loaded_at=now_utc(),
        )
        normalizer = DocumentNormalizer()
        dataset = normalizer.normalize(raw)

        assert dataset.dataframe["section_type"].iloc[0] == "code"
        assert "print('hello')" in dataset.dataframe["content"].iloc[0]

    def test_rejects_tabular_data(self):
        raw = RawData(
            id=generate_id(),
            source_config=SourceConfig(
                source_type=SourceType.CSV,
                name="tabular",
                file=FileSourceConfig(path=Path("/test/data.csv")),
            ),
            content_type=ContentType.TABULAR,
            tabular_data=[{"a": "1"}],
            tabular_headers=["a"],
            row_count=1,
            loaded_at=now_utc(),
        )
        normalizer = DocumentNormalizer()
        with pytest.raises(NormalizationError, match="Expected document"):
            normalizer.normalize(raw)

    def test_empty_document(self):
        raw = RawData(
            id=generate_id(),
            source_config=SourceConfig(
                source_type=SourceType.TXT,
                name="empty",
                file=FileSourceConfig(path=Path("/test/empty.txt")),
            ),
            content_type=ContentType.DOCUMENT,
            document_text=[],
            loaded_at=now_utc(),
        )
        normalizer = DocumentNormalizer()
        dataset = normalizer.normalize(raw)
        assert dataset.row_count == 0

    def test_provenance_tracked(self):
        raw = RawData(
            id=generate_id(),
            source_config=SourceConfig(
                source_type=SourceType.DOCX,
                name="my_report",
                file=FileSourceConfig(path=Path("/test/report.docx")),
            ),
            content_type=ContentType.DOCUMENT,
            document_text=["Some text."],
            loaded_at=now_utc(),
        )
        normalizer = DocumentNormalizer()
        dataset = normalizer.normalize(raw)

        assert dataset.provenance.source_type == "docx"
        assert dataset.provenance.source_name == "my_report"
        assert dataset.provenance.source_path == "/test/report.docx"


# ---- Default Registry ----


class TestDefaultRegistryAllTypes:
    def test_all_types_registered(self):
        registry = create_default_registry()
        types = registry.available_types()

        assert "csv" in types
        assert "json" in types
        assert "xlsx" in types
        assert "docx" in types
        assert "txt" in types
        assert "md" in types

    def test_get_each_loader_type(self):
        registry = create_default_registry()

        for source_type in (
            SourceType.CSV,
            SourceType.JSON,
            SourceType.XLSX,
            SourceType.DOCX,
            SourceType.TXT,
            SourceType.MD,
        ):
            loader = registry.get_loader(source_type)
            assert loader is not None
