"""Tests for DOCX and XLSX output adapters and output registry (Phase 15)."""

from __future__ import annotations

import io
from datetime import UTC, datetime

import docx
import pytest
from openpyxl import load_workbook

from pygramattic_reports.models import Report, ReportSection, SectionType, ThemeSpec
from pygramattic_reports.outputs import (
    DocxAdapter,
    MarkdownAdapter,
    XlsxAdapter,
    create_default_output_registry,
)

# ---------- Helpers ----------------------------------------------------------


def _make_report(
    sections: list[ReportSection],
    name: str = "Test Report",
) -> Report:
    return Report(
        id="rpt-001",
        name=name,
        sections=sections,
        template_name="test",
        theme_name="test",
        build_timestamp=datetime(2025, 1, 1, tzinfo=UTC),
    )


def _default_theme() -> ThemeSpec:
    return ThemeSpec(name="test")


def _open_docx(data: bytes) -> docx.document.Document:
    return docx.Document(io.BytesIO(data))


# ---------- DOCX: Title ------------------------------------------------------


class TestDocxTitle:
    """Title renders as heading level 0."""

    def test_docx_title(self) -> None:
        report = _make_report(
            [
                ReportSection(section_type=SectionType.TITLE, content="My Report"),
            ]
        )
        data = DocxAdapter().render(report, _default_theme())
        doc = _open_docx(data)
        assert any(
            p.text == "My Report" and p.style is not None and p.style.name.startswith("Title")
            for p in doc.paragraphs
        )


# ---------- DOCX: Headings ---------------------------------------------------


class TestDocxHeadings:
    """Headings use correct levels."""

    def test_docx_heading_level_2(self) -> None:
        report = _make_report(
            [
                ReportSection(
                    section_type=SectionType.HEADING,
                    content="Section",
                    level=2,
                ),
            ]
        )
        data = DocxAdapter().render(report, _default_theme())
        doc = _open_docx(data)
        assert any(
            p.text == "Section" and p.style is not None and "Heading 2" in p.style.name
            for p in doc.paragraphs
        )

    def test_docx_heading_clamped_to_4(self) -> None:
        report = _make_report(
            [
                ReportSection(
                    section_type=SectionType.HEADING,
                    content="Deep",
                    level=10,
                ),
            ]
        )
        data = DocxAdapter().render(report, _default_theme())
        doc = _open_docx(data)
        assert any(
            p.text == "Deep" and p.style is not None and "Heading 4" in p.style.name
            for p in doc.paragraphs
        )


# ---------- DOCX: Narrative --------------------------------------------------


class TestDocxNarrative:
    """Text renders as paragraphs."""

    def test_docx_narrative(self) -> None:
        report = _make_report(
            [
                ReportSection(
                    section_type=SectionType.NARRATIVE,
                    content="First paragraph.\n\nSecond paragraph.",
                ),
            ]
        )
        data = DocxAdapter().render(report, _default_theme())
        doc = _open_docx(data)
        texts = [p.text for p in doc.paragraphs]
        assert "First paragraph." in texts
        assert "Second paragraph." in texts


# ---------- DOCX: Data table --------------------------------------------------


class TestDocxDataTable:
    """Table renders with headers and data."""

    def test_docx_data_table(self) -> None:
        report = _make_report(
            [
                ReportSection(
                    section_type=SectionType.DATA_TABLE,
                    title="Revenue",
                    table_data={
                        "headers": ["Region", "Amount"],
                        "rows": [["US", 1500], ["EU", 2300]],
                    },
                ),
            ]
        )
        data = DocxAdapter().render(report, _default_theme())
        doc = _open_docx(data)
        assert len(doc.tables) == 1
        table = doc.tables[0]
        assert table.rows[0].cells[0].text == "Region"
        assert table.rows[0].cells[1].text == "Amount"
        assert table.rows[1].cells[0].text == "US"
        assert table.rows[1].cells[1].text == "1500"
        assert table.rows[2].cells[0].text == "EU"


# ---------- DOCX: Chart embedded ---------------------------------------------


class TestDocxChartEmbedded:
    """Chart image is embedded."""

    def test_docx_chart_embedded(self) -> None:
        report = _make_report(
            [
                ReportSection(
                    section_type=SectionType.CHART,
                    title="Chart",
                    media_bytes=_minimal_png(),
                    media_type="image/png",
                ),
            ]
        )
        data = DocxAdapter().render(report, _default_theme())
        doc = _open_docx(data)
        inline_shapes = doc.inline_shapes
        assert len(inline_shapes) == 1


# ---------- DOCX: Page break -------------------------------------------------


class TestDocxPageBreak:
    """Page break inserts correctly."""

    def test_docx_page_break(self) -> None:
        report = _make_report(
            [
                ReportSection(section_type=SectionType.TITLE, content="Before"),
                ReportSection(section_type=SectionType.PAGE_BREAK),
                ReportSection(section_type=SectionType.TITLE, content="After"),
            ]
        )
        data = DocxAdapter().render(report, _default_theme())
        assert len(data) > 0
        doc = _open_docx(data)
        texts = [p.text for p in doc.paragraphs]
        assert "Before" in texts
        assert "After" in texts


# ---------- DOCX: Valid file (roundtrip) -------------------------------------


class TestDocxValidFile:
    """Output opens with python-docx (roundtrip test)."""

    def test_docx_valid_file(self) -> None:
        report = _make_report(
            [
                ReportSection(section_type=SectionType.TITLE, content="Test"),
                ReportSection(
                    section_type=SectionType.NARRATIVE,
                    content="Body text.",
                ),
                ReportSection(
                    section_type=SectionType.DATA_TABLE,
                    table_data={
                        "headers": ["A"],
                        "rows": [["x"]],
                    },
                ),
            ]
        )
        data = DocxAdapter().render(report, _default_theme())
        doc = _open_docx(data)
        assert len(doc.paragraphs) > 0


# ---------- DOCX: Theme styling ----------------------------------------------


class TestDocxThemeStyling:
    """Theme fonts and colors applied."""

    def test_docx_theme_styling(self) -> None:
        report = _make_report(
            [
                ReportSection(section_type=SectionType.TITLE, content="Styled"),
            ]
        )
        data = DocxAdapter().render(report, _default_theme())
        doc = _open_docx(data)
        title_para = next(p for p in doc.paragraphs if p.text == "Styled")
        assert len(title_para.runs) > 0
        run = title_para.runs[0]
        assert run.font.name == "Arial"


# ---------- XLSX: Data table --------------------------------------------------


class TestXlsxDataTable:
    """XLSX contains correct data."""

    def test_xlsx_data_table(self) -> None:
        report = _make_report(
            [
                ReportSection(
                    section_type=SectionType.DATA_TABLE,
                    title="Sales",
                    table_data={
                        "headers": ["Product", "Qty"],
                        "rows": [["Widget", 100], ["Gadget", 200]],
                    },
                ),
            ]
        )
        data = XlsxAdapter().render(report, _default_theme())
        wb = load_workbook(io.BytesIO(data))
        ws = wb.active
        assert ws is not None
        assert ws["A1"].value == "Product"
        assert ws["B1"].value == "Qty"
        assert ws["A2"].value == "Widget"
        assert ws["B2"].value == 100
        assert ws["A3"].value == "Gadget"


# ---------- XLSX: Header styling ---------------------------------------------


class TestXlsxHeaderStyling:
    """Headers are styled."""

    def test_xlsx_header_styling(self) -> None:
        report = _make_report(
            [
                ReportSection(
                    section_type=SectionType.DATA_TABLE,
                    title="Data",
                    table_data={
                        "headers": ["Col"],
                        "rows": [["val"]],
                    },
                ),
            ]
        )
        data = XlsxAdapter().render(report, _default_theme())
        wb = load_workbook(io.BytesIO(data))
        ws = wb.active
        assert ws is not None
        header_cell = ws["A1"]
        assert header_cell.font.bold is True


# ---------- XLSX: Multiple tables → multiple sheets --------------------------


class TestXlsxMultipleTables:
    """Multiple tables create multiple sheets."""

    def test_xlsx_multiple_tables(self) -> None:
        report = _make_report(
            [
                ReportSection(
                    section_type=SectionType.DATA_TABLE,
                    title="First",
                    table_data={
                        "headers": ["A"],
                        "rows": [["1"]],
                    },
                ),
                ReportSection(
                    section_type=SectionType.DATA_TABLE,
                    title="Second",
                    table_data={
                        "headers": ["B"],
                        "rows": [["2"]],
                    },
                ),
            ]
        )
        data = XlsxAdapter().render(report, _default_theme())
        wb = load_workbook(io.BytesIO(data))
        assert len(wb.sheetnames) == 2
        assert wb.sheetnames[0] == "First"
        assert wb.sheetnames[1] == "Second"


# ---------- XLSX: No tables ---------------------------------------------------


class TestXlsxNoTables:
    """Report with no tables gets placeholder."""

    def test_xlsx_no_tables(self) -> None:
        report = _make_report(
            [
                ReportSection(section_type=SectionType.TITLE, content="No data"),
            ]
        )
        data = XlsxAdapter().render(report, _default_theme())
        wb = load_workbook(io.BytesIO(data))
        ws = wb.active
        assert ws is not None
        assert ws["A1"].value == "No data tables in this report"


# ---------- Output registry ---------------------------------------------------


class TestOutputRegistry:
    """Registry returns correct adapters."""

    def test_output_registry(self) -> None:
        registry = create_default_output_registry()
        assert isinstance(registry.get_adapter("md"), MarkdownAdapter)
        assert isinstance(registry.get_adapter("markdown"), MarkdownAdapter)
        assert isinstance(registry.get_adapter("docx"), DocxAdapter)
        assert isinstance(registry.get_adapter("xlsx"), XlsxAdapter)

    def test_registry_available_formats(self) -> None:
        registry = create_default_output_registry()
        formats = registry.available_formats()
        assert "md" in formats
        assert "docx" in formats
        assert "xlsx" in formats

    def test_registry_unknown_raises(self) -> None:
        registry = create_default_output_registry()
        with pytest.raises(ValueError, match="pdf"):
            registry.get_adapter("pdf")


# ---------- Helpers ----------------------------------------------------------


def _minimal_png() -> bytes:
    """Return a minimal valid PNG for embedding tests."""
    import struct  # noqa: PLC0415
    import zlib  # noqa: PLC0415

    def _chunk(chunk_type: bytes, data: bytes) -> bytes:
        raw = chunk_type + data
        return struct.pack(">I", len(data)) + raw + struct.pack(">I", zlib.crc32(raw) & 0xFFFFFFFF)

    sig = b"\x89PNG\r\n\x1a\n"
    ihdr = _chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0))
    raw_data = zlib.compress(b"\x00\xff\xff\xff")
    idat = _chunk(b"IDAT", raw_data)
    iend = _chunk(b"IEND", b"")
    return sig + ihdr + idat + iend
