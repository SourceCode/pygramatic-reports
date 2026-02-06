"""DOCX (Word) file loader."""

from __future__ import annotations

from typing import TYPE_CHECKING, Any

import docx

from pygramattic_reports.exceptions import LoaderError
from pygramattic_reports.logging import get_logger
from pygramattic_reports.models import ContentType, RawData, SourceType
from pygramattic_reports.models.base import generate_id, now_utc
from pygramattic_reports.storage.checksum import compute_checksum

from .base import BaseLoader

if TYPE_CHECKING:
    from pathlib import Path

    from pygramattic_reports.models import SourceConfig

_logger = get_logger("loaders.docx")

_HEADING_STYLES: dict[str, int] = {
    "Heading 1": 1,
    "Heading 2": 2,
    "Heading 3": 3,
    "Heading 4": 4,
    "Heading 5": 5,
    "Heading 6": 6,
    "Title": 1,
}


class DocxLoader(BaseLoader):
    """Loads Word documents into RawData objects.

    Uses ``python-docx`` to extract paragraphs and tables.
    Heading styles are converted to markdown-style ``#`` prefixes.
    """

    def load(self, config: SourceConfig) -> RawData:
        """Load a DOCX file into a RawData object.

        Args:
            config: Source configuration with file path and options.

        Returns:
            RawData with ``content_type=DOCUMENT``.

        Raises:
            LoaderError: If the file cannot be read or parsed.
        """
        if config.file is None:
            msg = "DOCX loader requires file configuration"
            raise LoaderError(msg, source_type="docx")

        file_path = config.file.path
        if not file_path.is_file():
            msg = f"File not found: {file_path}"
            raise LoaderError(msg, source_type="docx", source_path=str(file_path))

        checksum = _compute_checksum(file_path)
        document = _open_document(file_path)

        document_text = _extract_paragraphs(document)
        tabular_data, tabular_headers = _extract_first_table(document)

        _logger.info(
            "Loaded DOCX file",
            path=str(file_path),
            paragraphs=len(document_text),
            has_table=tabular_headers is not None,
        )

        return RawData(
            id=generate_id(),
            source_config=config,
            content_type=ContentType.DOCUMENT,
            document_text=document_text,
            tabular_data=tabular_data,
            tabular_headers=tabular_headers,
            loaded_at=now_utc(),
            source_checksum=checksum,
        )

    def supported_types(self) -> list[str]:
        """Return supported source types."""
        return [SourceType.DOCX]


def _compute_checksum(file_path: Path) -> str:
    """Compute SHA-256 checksum of the source file."""
    try:
        return compute_checksum(file_path)
    except OSError as exc:
        msg = f"Failed to compute checksum for {file_path}: {exc}"
        raise LoaderError(msg, source_type="docx", source_path=str(file_path)) from exc


def _open_document(file_path: Path) -> Any:  # noqa: ANN401
    """Open a DOCX file and return the Document object."""
    try:
        return docx.Document(str(file_path))
    except Exception as exc:
        msg = f"Failed to open DOCX {file_path}: {exc}"
        raise LoaderError(msg, source_type="docx", source_path=str(file_path)) from exc


def _extract_paragraphs(document: Any) -> list[str]:  # noqa: ANN401
    """Extract paragraph text, converting heading styles to # prefixes."""
    blocks: list[str] = []
    for para in document.paragraphs:
        text: str = para.text.strip()
        if not text:
            continue
        style_name: str = para.style.name if para.style else ""
        level = _HEADING_STYLES.get(style_name, 0)
        if level > 0:
            blocks.append("#" * level + " " + text)
        else:
            blocks.append(text)
    return blocks


def _extract_first_table(
    document: Any,  # noqa: ANN401
) -> tuple[list[dict[str, Any]] | None, list[str] | None]:
    """Extract the first table from the document, if any."""
    if not document.tables:
        return None, None

    table = document.tables[0]
    rows = table.rows
    if len(rows) < 1:
        return None, None

    headers: list[str] = [cell.text.strip() for cell in rows[0].cells]
    tabular_data: list[dict[str, Any]] = []
    for row in rows[1:]:
        values: list[str] = [cell.text.strip() for cell in row.cells]
        tabular_data.append(dict(zip(headers, values, strict=True)))

    return tabular_data, headers
