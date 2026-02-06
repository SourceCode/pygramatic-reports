"""XLSX (Excel) file loader."""

from __future__ import annotations

from datetime import datetime
from typing import TYPE_CHECKING, Any

import openpyxl

from pygramattic_reports.exceptions import LoaderError
from pygramattic_reports.logging import get_logger
from pygramattic_reports.models import ContentType, RawData, SourceType
from pygramattic_reports.models.base import generate_id, now_utc
from pygramattic_reports.storage.checksum import compute_checksum

from .base import BaseLoader

if TYPE_CHECKING:
    from pathlib import Path

    from openpyxl.worksheet.worksheet import Worksheet

    from pygramattic_reports.models import SourceConfig

_logger = get_logger("loaders.xlsx")


class XlsxLoader(BaseLoader):
    """Loads Excel workbooks into RawData objects.

    Uses ``openpyxl`` to read ``.xlsx`` files. Handles sheet selection,
    merged cells, header detection, and date formatting.
    """

    def load(self, config: SourceConfig) -> RawData:
        """Load an XLSX file into a RawData object.

        Args:
            config: Source configuration with file path and options.

        Returns:
            RawData with ``content_type=TABULAR``.

        Raises:
            LoaderError: If the file cannot be read or parsed.
        """
        if config.file is None:
            msg = "XLSX loader requires file configuration"
            raise LoaderError(msg, source_type="xlsx")

        file_path = config.file.path
        if not file_path.is_file():
            msg = f"File not found: {file_path}"
            raise LoaderError(msg, source_type="xlsx", source_path=str(file_path))

        checksum = _compute_checksum(file_path)
        ws = _open_worksheet(file_path, config.file.sheet_name)
        _fill_merged_cells(ws)

        headers, tabular_data = _read_sheet(ws, has_header=config.file.has_header)

        _logger.info(
            "Loaded XLSX file",
            path=str(file_path),
            sheet=ws.title,
            rows=len(tabular_data),
            columns=len(headers),
        )

        return RawData(
            id=generate_id(),
            source_config=config,
            content_type=ContentType.TABULAR,
            tabular_data=tabular_data,
            tabular_headers=headers,
            row_count=len(tabular_data),
            loaded_at=now_utc(),
            source_checksum=checksum,
        )

    def supported_types(self) -> list[str]:
        """Return supported source types."""
        return [SourceType.XLSX]


def _compute_checksum(file_path: Path) -> str:
    """Compute SHA-256 checksum of the source file."""
    try:
        return compute_checksum(file_path)
    except OSError as exc:
        msg = f"Failed to compute checksum for {file_path}: {exc}"
        raise LoaderError(msg, source_type="xlsx", source_path=str(file_path)) from exc


def _open_worksheet(file_path: Path, sheet_name: str | None) -> Worksheet:
    """Open the workbook and return the target worksheet."""
    try:
        wb = openpyxl.load_workbook(file_path, data_only=True, read_only=False)
    except Exception as exc:
        msg = f"Failed to open workbook {file_path}: {exc}"
        raise LoaderError(msg, source_type="xlsx", source_path=str(file_path)) from exc

    if not wb.sheetnames:
        msg = f"Workbook has no sheets: {file_path}"
        raise LoaderError(msg, source_type="xlsx", source_path=str(file_path))

    if sheet_name is not None:
        if sheet_name not in wb.sheetnames:
            msg = f"Sheet {sheet_name!r} not found in {file_path} (available: {wb.sheetnames})"
            raise LoaderError(msg, source_type="xlsx", source_path=str(file_path))
        return wb[sheet_name]

    return wb.active  # type: ignore[return-value]


def _fill_merged_cells(ws: Worksheet) -> None:
    """Unmerge cells and fill values into the merged range."""
    for merge_range in list(ws.merged_cells.ranges):
        value = ws.cell(merge_range.min_row, merge_range.min_col).value
        ws.unmerge_cells(str(merge_range))
        for row in range(merge_range.min_row, merge_range.max_row + 1):
            for col in range(merge_range.min_col, merge_range.max_col + 1):
                ws.cell(row=row, column=col, value=value)


def _read_sheet(
    ws: Worksheet,
    *,
    has_header: bool,
) -> tuple[list[str], list[dict[str, Any]]]:
    """Read rows from a worksheet, returning headers and row dicts."""
    rows_iter = ws.iter_rows(values_only=True)

    if has_header:
        header_row = next(rows_iter, None)
        if header_row is None:
            return [], []
        headers = [
            str(cell) if cell is not None else f"col_{i}" for i, cell in enumerate(header_row)
        ]
    else:
        first_row = next(rows_iter, None)
        if first_row is None:
            return [], []
        headers = [f"col_{i}" for i in range(len(first_row))]
        converted = [_convert_cell(v) for v in first_row]
        if not _is_empty(converted):
            tabular_data = [dict(zip(headers, converted, strict=True))]
        else:
            tabular_data = []
        for row in rows_iter:
            values = [_convert_cell(v) for v in row]
            if _is_empty(values):
                continue
            tabular_data.append(dict(zip(headers, values, strict=True)))
        return headers, tabular_data

    tabular_data = []
    for row in rows_iter:
        values = [_convert_cell(v) for v in row]
        if _is_empty(values):
            continue
        tabular_data.append(dict(zip(headers, values, strict=True)))

    return headers, tabular_data


def _convert_cell(value: object) -> Any:  # noqa: ANN401
    """Convert an openpyxl cell value to a serializable Python type."""
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat()
    return value


def _is_empty(values: list[Any]) -> bool:
    """Check if all values in a row are None."""
    return all(v is None for v in values)
