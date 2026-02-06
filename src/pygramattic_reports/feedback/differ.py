"""Report differ for comparing original and edited reports.

Produces structured diffs for Markdown, TXT, and DOCX formats.
"""

from __future__ import annotations

import difflib
from dataclasses import dataclass, field
from typing import TYPE_CHECKING

from pygramattic_reports.exceptions import FeedbackError

if TYPE_CHECKING:
    from pathlib import Path


@dataclass
class DiffEntry:
    """A single difference between original and edited report.

    Attributes:
        section: Section identifier or line range.
        change_type: One of ``"added"``, ``"removed"``, ``"modified"``.
        original: Original text (empty for added).
        edited: Edited text (empty for removed).
    """

    section: str
    change_type: str
    original: str
    edited: str


@dataclass
class ReportDiff:
    """Structured diff between original and edited report.

    Attributes:
        original_path: Path to the original file.
        edited_path: Path to the edited file.
        format: File format (``"markdown"``, ``"text"``, ``"docx"``).
        entries: List of diff entries.
        total_changes: Number of changes found.
        summary: Human-readable summary of changes.
    """

    original_path: str
    edited_path: str
    format: str
    entries: list[DiffEntry] = field(default_factory=list)
    total_changes: int = 0
    summary: str = "No changes detected"


class ReportDiffer:
    """Compares original and edited reports to identify changes.

    Supports Markdown, TXT, and DOCX comparison. Produces a structured
    diff that the feedback analyzer can use to suggest adjustments.

    Usage::

        differ = ReportDiffer()
        diff = differ.compare(
            original=Path("report_v1.md"),
            edited=Path("report_v1_edited.md"),
        )
    """

    def compare(self, original: Path, edited: Path) -> ReportDiff:
        """Compare two report files.

        Auto-detects format from file extension.

        Args:
            original: Path to the original generated report.
            edited: Path to the user-edited version.

        Returns:
            Structured diff of changes.

        Raises:
            FeedbackError: If files cannot be compared.
        """
        if not original.exists():
            raise FeedbackError(f"Original file not found: {original}")
        if not edited.exists():
            raise FeedbackError(f"Edited file not found: {edited}")

        ext = original.suffix.lower()
        if ext == ".md":
            return self._compare_text(original, edited, "markdown")
        if ext == ".txt":
            return self._compare_text(original, edited, "text")
        if ext == ".docx":
            return self._compare_docx(original, edited)
        raise FeedbackError(f"Unsupported format for comparison: {ext}")

    def _compare_text(
        self,
        original: Path,
        edited: Path,
        fmt: str,
    ) -> ReportDiff:
        """Compare two text-based files."""
        orig_lines = original.read_text(encoding="utf-8").splitlines()
        edit_lines = edited.read_text(encoding="utf-8").splitlines()

        diff = difflib.unified_diff(orig_lines, edit_lines, lineterm="")
        entries = self._parse_unified_diff(list(diff))

        return ReportDiff(
            original_path=str(original),
            edited_path=str(edited),
            format=fmt,
            entries=entries,
            total_changes=len(entries),
            summary=self._generate_summary(entries),
        )

    def _compare_docx(
        self,
        original: Path,
        edited: Path,
    ) -> ReportDiff:
        """Compare two DOCX files by extracting text and diffing."""
        import docx  # noqa: PLC0415

        orig_doc = docx.Document(str(original))
        edit_doc = docx.Document(str(edited))

        orig_text = [p.text for p in orig_doc.paragraphs if p.text.strip()]
        edit_text = [p.text for p in edit_doc.paragraphs if p.text.strip()]

        diff = difflib.unified_diff(orig_text, edit_text, lineterm="")
        entries = self._parse_unified_diff(list(diff))

        return ReportDiff(
            original_path=str(original),
            edited_path=str(edited),
            format="docx",
            entries=entries,
            total_changes=len(entries),
            summary=self._generate_summary(entries),
        )

    @staticmethod
    def _parse_unified_diff(diff_lines: list[str]) -> list[DiffEntry]:
        """Parse unified diff output into structured DiffEntry objects."""
        entries: list[DiffEntry] = []
        current_section = "header"
        removed_buf: list[str] = []
        added_buf: list[str] = []

        def _flush() -> None:
            """Flush buffered additions/removals into entries."""
            if removed_buf and added_buf:
                entries.append(
                    DiffEntry(
                        section=current_section,
                        change_type="modified",
                        original="\n".join(removed_buf),
                        edited="\n".join(added_buf),
                    )
                )
            elif removed_buf:
                entries.append(
                    DiffEntry(
                        section=current_section,
                        change_type="removed",
                        original="\n".join(removed_buf),
                        edited="",
                    )
                )
            elif added_buf:
                entries.append(
                    DiffEntry(
                        section=current_section,
                        change_type="added",
                        original="",
                        edited="\n".join(added_buf),
                    )
                )
            removed_buf.clear()
            added_buf.clear()

        for line in diff_lines:
            if line.startswith("@@"):
                _flush()
                current_section = line
            elif line.startswith("-") and not line.startswith("---"):
                removed_buf.append(line[1:])
            elif line.startswith("+") and not line.startswith("+++"):
                added_buf.append(line[1:])
            elif removed_buf or added_buf:
                _flush()

        _flush()
        return entries

    @staticmethod
    def _generate_summary(entries: list[DiffEntry]) -> str:
        """Generate a human-readable summary of changes."""
        added = sum(1 for e in entries if e.change_type == "added")
        removed = sum(1 for e in entries if e.change_type == "removed")
        modified = sum(1 for e in entries if e.change_type == "modified")
        parts: list[str] = []
        if added:
            parts.append(f"{added} additions")
        if removed:
            parts.append(f"{removed} removals")
        if modified:
            parts.append(f"{modified} modifications")
        return ", ".join(parts) if parts else "No changes detected"
