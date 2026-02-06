"""PowerPoint output adapter for pygramattic-reports.

Generates PPTX files using python-pptx.
"""

from __future__ import annotations

from typing import TYPE_CHECKING, Any

try:
    import pptx
    from pptx.util import Inches, Pt

    HAS_PPTX = True
except ImportError:
    HAS_PPTX = False

from pygramattic_reports.logging import get_logger
from pygramattic_reports.models import SectionType
from pygramattic_reports.outputs.base import BaseOutputAdapter

if TYPE_CHECKING:
    from pygramattic_reports.models import Report, ReportSection, ThemeSpec

logger = get_logger("outputs.pptx")


class PptxAdapter(BaseOutputAdapter):
    """Generates PowerPoint presentations."""

    def __init__(self) -> None:
        """Initialize and check dependencies."""
        if not HAS_PPTX:
            raise ImportError(
                "python-pptx is required for PowerPoint output. "
                "Install it with: pip install python-pptx"
            )

    def render(self, report: Report, theme: ThemeSpec) -> bytes:
        """Render report to PPTX bytes."""
        prs = pptx.Presentation()

        # 1. Title Slide
        self._render_title_slide(prs, report)

        # 2. Render Sections
        for section in report.sections:
            if section.section_type in (SectionType.TITLE, SectionType.HEADING):
                # Major section break / title slide
                self._render_section_header_slide(prs, section)

            elif (
                section.section_type
                in (SectionType.NARRATIVE, SectionType.SUMMARY, SectionType.CALLOUT)
                or section.section_type == SectionType.LIST
            ):
                self._render_content_slide(prs, section)

            elif section.section_type == SectionType.CHART:
                self._render_chart_slide(prs, section)

            elif section.section_type == SectionType.DATA_TABLE:
                # Tables in PPT are tricky to do well automatically without overflowing
                # For now, render as a content slide (if small) or skip?
                # Let's try to render a simple table
                self._render_table_slide(prs, section)

        # Save to bytes
        from io import BytesIO

        output = BytesIO()
        prs.save(output)
        return output.getvalue()

    def file_extension(self) -> str:
        return ".pptx"

    def _render_title_slide(self, prs: Any, report: Report) -> None:
        """Create the main title slide."""
        # layout 0 is usually Title Slide
        slide_layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(slide_layout)

        title = slide.shapes.title
        subtitle = slide.placeholders[1]

        title.text = report.name

        lines = []
        if report.metadata and report.metadata.author:
            lines.append(f"Author: {report.metadata.author}")
        lines.append(f"Date: {report.build_timestamp.strftime('%Y-%m-%d')}")

        subtitle.text = "\n".join(lines)

    def _render_section_header_slide(self, prs: Any, section: ReportSection) -> None:
        """Create a section header slide."""
        # layout 2 is usually Section Header
        slide_layout = prs.slide_layouts[2]
        slide = prs.slides.add_slide(slide_layout)

        title = slide.shapes.title
        title.text = section.title or "Section"

        # Add content as subtitle if present
        if section.content:
            try:
                subtitle = slide.placeholders[1]
                subtitle.text = section.content
            except IndexError:
                pass

    def _render_content_slide(self, prs: Any, section: ReportSection) -> None:
        """Create a standard title + content slide."""
        # layout 1 is usually Title and Content
        slide_layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(slide_layout)

        title = slide.shapes.title
        title.text = section.title or section.section_type.title()

        body = slide.placeholders[1]
        tf = body.text_frame

        if section.content:
            tf.text = section.content

        elif section.list_data:
            tf.clear()  # Clear default empty para
            for item in section.list_data:
                p = tf.add_paragraph()
                p.text = str(item)
                p.level = 0

        elif section.callout_data:
            tf.text = section.callout_data.get("content", "")

    def _render_chart_slide(self, prs: Any, section: ReportSection) -> None:
        """Create a slide with a logical chart image."""
        if not section.media_bytes:
            return

        # Use a blank layout or Title Only
        slide_layout = prs.slide_layouts[5]  # Title Only
        slide = prs.slides.add_slide(slide_layout)

        title = slide.shapes.title
        title.text = section.title or "Chart"

        from io import BytesIO

        img_stream = BytesIO(section.media_bytes)

        # Center the image
        # Slide width is usually 10 inches, height 7.5
        # Leave space for title

        left = Inches(1)
        top = Inches(1.5)
        width = Inches(8)
        height = Inches(5.5)

        slide.shapes.add_picture(img_stream, left, top, width=width, height=height)

    def _render_table_slide(self, prs: Any, section: ReportSection) -> None:
        """Create a slide with a table."""
        if not section.table_data:
            return

        headers = section.table_data.get("headers", [])
        rows = section.table_data.get("rows", [])

        if not headers:
            return

        # Cap rows to fit on one slide
        max_rows = 10
        display_rows = rows[:max_rows]

        slide_layout = prs.slide_layouts[5]  # Title Only
        slide = prs.slides.add_slide(slide_layout)
        title = slide.shapes.title
        title.text = section.title or "Data Table"

        rows_count = len(display_rows) + 1  # + header
        cols_count = len(headers)

        left = Inches(0.5)
        top = Inches(1.5)
        width = Inches(9)
        height = Inches(0.5 * rows_count)

        table = slide.shapes.add_table(rows_count, cols_count, left, top, width, height).table

        # headers
        for i, h in enumerate(headers):
            table.cell(0, i).text = str(h)

        # rows
        for r_idx, row in enumerate(display_rows):
            for c_idx, val in enumerate(row):
                if c_idx < cols_count:
                    table.cell(r_idx + 1, c_idx).text = str(val)
