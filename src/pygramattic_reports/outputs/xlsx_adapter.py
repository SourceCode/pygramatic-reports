"""XLSX (Excel) output adapter for pygramattic-reports.

Exports data tables from reports as styled Excel spreadsheets
using ``openpyxl``.
"""

from __future__ import annotations

import io
from typing import TYPE_CHECKING, Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from pygramattic_reports.models import SectionType

from .base import BaseOutputAdapter

if TYPE_CHECKING:
    from openpyxl.worksheet.worksheet import Worksheet

    from pygramattic_reports.models import Report, ReportSection, ThemeSpec


class XlsxAdapter(BaseOutputAdapter):
    """Renders report data tables as Excel spreadsheets.

    Creates one worksheet per data table section in the report.
    Applies theme-driven header styling and writes data rows.
    """

    def render(self, report: Report, theme: ThemeSpec) -> bytes:
        """Render data tables from the report as XLSX.

        Args:
            report: The abstract report to render.
            theme: Theme for header styling.

        Returns:
            XLSX file bytes.
        """
        wb = Workbook()
        ws: Any = wb.active

        tables = [
            s for s in report.sections
            if s.section_type == SectionType.DATA_TABLE
        ]

        if not tables:
            ws.title = "No Data"
            ws["A1"] = "No data tables in this report"
        else:
            for i, table_section in enumerate(tables):
                if i == 0:
                    ws.title = (table_section.title or "Data")[:31]
                else:
                    ws = wb.create_sheet(
                        title=(table_section.title or f"Sheet{i + 1}")[:31],
                    )
                self._write_table(ws, table_section, theme)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf.read()

    def file_extension(self) -> str:
        """Return the XLSX file extension."""
        return ".xlsx"

    @staticmethod
    def _write_table(
        ws: Worksheet,
        section: ReportSection,
        theme: ThemeSpec,
    ) -> None:
        """Write a data table to a worksheet with styled headers."""
        if not section.table_data:
            return

        headers: list[str] = section.table_data["headers"]
        rows: list[list[object]] = section.table_data["rows"]

        primary = theme.colors.primary.lstrip("#")
        header_fill = PatternFill(
            start_color=primary,
            end_color=primary,
            fill_type="solid",
        )
        header_font = Font(bold=True, color="FFFFFF")

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font

        for row_idx, row_data in enumerate(rows, 2):
            for col_idx, value in enumerate(row_data, 1):
                ws.cell(row=row_idx, column=col_idx, value=value)  # type: ignore[call-overload]
