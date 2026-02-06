"""DOCX (Word) output adapter for pygramattic-reports.

Converts an abstract Report into a professionally styled Word document
using ``python-docx``.
"""

from __future__ import annotations

import io
from typing import TYPE_CHECKING, Any

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from pygramattic_reports.models import SectionType
from pygramattic_reports.themes import ThemeApplicator

from .base import BaseOutputAdapter

if TYPE_CHECKING:
    from pygramattic_reports.models import Report, ReportSection, ThemeSpec


def _hex_to_rgb(hex_color: str) -> RGBColor:
    """Convert a hex color string to a ``docx.shared.RGBColor``."""
    h = hex_color.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


class DocxAdapter(BaseOutputAdapter):
    """Renders reports as DOCX (Word) documents.

    Produces professional Word documents with:

    - Themed heading styles
    - Formatted data tables with header styling
    - Embedded chart images
    - Consistent font and color theming
    - Proper page margins
    """

    def render(self, report: Report, theme: ThemeSpec) -> bytes:
        """Render the report as a DOCX document.

        Args:
            report: The abstract report to render.
            theme: Theme for styling.

        Returns:
            DOCX file bytes.
        """
        doc = docx.Document()
        applicator = ThemeApplicator(theme)
        styles = applicator.to_docx_styles()

        self._setup_page(doc, styles)

        for section in report.sections:
            self._render_section(doc, section, styles)

        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        return buf.read()

    def file_extension(self) -> str:
        """Return the DOCX file extension."""
        return ".docx"

    # -- Internal helpers -----------------------------------------------------

    @staticmethod
    def _setup_page(
        doc: Any,  # noqa: ANN401
        styles: dict[str, Any],
    ) -> None:
        """Configure page margins from theme spacing."""
        margin = Inches(styles["spacing"]["page_margin_inches"])
        page = doc.sections[0]
        page.top_margin = margin
        page.bottom_margin = margin
        page.left_margin = margin
        page.right_margin = margin

    def _render_section(
        self,
        doc: Any,  # noqa: ANN401
        section: ReportSection,
        styles: dict[str, Any],
    ) -> None:
        """Dispatch section rendering by type."""
        st = section.section_type

        if st in (SectionType.TITLE, SectionType.HEADING):
            self._add_heading(doc, section, styles)
        elif st in (SectionType.SUMMARY, SectionType.NARRATIVE):
            self._add_text(doc, section, styles)
        elif st == SectionType.DATA_TABLE:
            self._add_table(doc, section, styles)
        elif st in (SectionType.CHART, SectionType.IMAGE):
            self._add_image(doc, section)
        elif st == SectionType.PAGE_BREAK:
            doc.add_page_break()
        elif st == SectionType.TABLE_OF_CONTENTS:
            doc.add_paragraph("[Table of Contents]")
        else:
            # SPACER or unknown
            doc.add_paragraph("")

    @staticmethod
    def _add_heading(
        doc: Any,  # noqa: ANN401
        section: ReportSection,
        styles: dict[str, Any],
    ) -> None:
        """Add a title or heading paragraph."""
        text = section.content or section.title or ""
        level = 0 if section.section_type == SectionType.TITLE else min(section.level, 4)

        para = doc.add_heading(text, level=level)
        for run in para.runs:
            run.font.name = styles["fonts"]["heading"]
            run.font.color.rgb = _hex_to_rgb(styles["colors"]["primary"])

    @staticmethod
    def _add_text(
        doc: Any,  # noqa: ANN401
        section: ReportSection,
        styles: dict[str, Any],
    ) -> None:
        """Add narrative or summary text."""
        if section.title:
            doc.add_heading(section.title, level=2)

        if section.content:
            paragraphs = section.content.split("\n\n")
            for para_text in paragraphs:
                para = doc.add_paragraph(para_text.strip())
                for run in para.runs:
                    run.font.name = styles["fonts"]["body"]
                    run.font.size = Pt(styles["sizes"]["body"])

    @staticmethod
    def _add_table(
        doc: Any,  # noqa: ANN401
        section: ReportSection,
        styles: dict[str, Any],
    ) -> None:
        """Add a data table with header styling."""
        if section.title:
            doc.add_heading(section.title, level=2)

        if not section.table_data:
            return

        headers: list[str] = section.table_data["headers"]
        rows: list[list[object]] = section.table_data["rows"]

        table = doc.add_table(rows=1 + len(rows), cols=len(headers))
        table.style = "Table Grid"

        header_rgb = _hex_to_rgb(styles["colors"]["primary"])
        body_font = styles["fonts"]["body"]
        body_size = Pt(styles["sizes"]["body"])

        for col_idx, header in enumerate(headers):
            cell = table.rows[0].cells[col_idx]
            cell.text = str(header)
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.bold = True
                    run.font.color.rgb = header_rgb
                    run.font.name = body_font
                    run.font.size = body_size

        for row_idx, row_data in enumerate(rows):
            for col_idx, value in enumerate(row_data):
                cell = table.rows[row_idx + 1].cells[col_idx]
                cell.text = str(value) if value is not None else ""

    @staticmethod
    def _add_image(
        doc: Any,  # noqa: ANN401
        section: ReportSection,
    ) -> None:
        """Add a chart or image."""
        if section.title:
            doc.add_heading(section.title, level=2)

        if section.media_bytes:
            stream = io.BytesIO(section.media_bytes)
            doc.add_picture(stream, width=Inches(6))
            last_para = doc.paragraphs[-1]
            last_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
